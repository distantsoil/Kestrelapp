"""File-format import: kestrel.docs.convert turns .md/.txt/.docx/.pdf/.doc into the markdown P7
decomposes. The property that matters is that the operator's wording survives verbatim (hard rule
2: imports._verbatim compares proposals against this text), that unsupported or oversize input is a
clear message rather than a crash, and that the web import route accepts an upload and shows what
it read before any proposal. Fixtures are fictional (Jordan Rivera, Northwind Freight, Cascade
Regional Health)."""

from __future__ import annotations

import asyncio
import io
import json
import os
import re
import stat
import sys
from pathlib import Path

import pytest

from kestrel.config import HostedMode as _HostedMode
from kestrel.docs import convert
from kestrel.docs.convert import (
    MAX_BYTES,
    ConvertError,
    ConvertResult,
    FileTooLarge,
    NoExtractableText,
    UnsupportedFormat,
    convert_to_markdown,
)
from kestrel.docs.imports import _verbatim

_TERMS = _HostedMode().disclosure_terms()

FIXTURES = Path(__file__).parent / "fixtures" / "docs"
BULLETS = [
    "Cut managed-service spend by a third across two regional distribution hubs.",
    "Led a HIPAA remediation that closed every audit finding in nine months.",
    "Replaced a legacy fleet-tracking platform with a cloud service in under a year.",
]


# ------------------------------------------------------------------------------ per format
def test_docx_converts_with_structure():
    r = convert_to_markdown(FIXTURES / "sample.docx")
    assert isinstance(r, ConvertResult)
    assert r.source_format == "docx"
    assert r.warnings == []
    assert "# Jordan Rivera" in r.markdown  # Heading 1 -> #
    assert "## Experience" in r.markdown  # Heading 2 -> ##
    assert "### Northwind Freight, VP of Information Technology (2019 to 2024)" in r.markdown
    assert "Cascade Regional Health" in r.markdown
    for b in BULLETS:
        assert f"- {b}" in r.markdown  # List Bullet style -> "- "
    assert "**Technology executive** with fifteen years" in r.markdown  # bold run -> **
    assert "| Security | Infrastructure |" in r.markdown  # table -> markdown table
    assert "| --- | --- |" in r.markdown
    assert "| HIPAA, SOC 2 | Azure, VMware |" in r.markdown
    # Paragraph order is preserved: name, summary, experience, then skills.
    assert r.markdown.index("# Jordan Rivera") < r.markdown.index("## Experience")
    assert r.markdown.index("Northwind") < r.markdown.index("Cascade")
    assert r.markdown.index("Cascade") < r.markdown.index("## Skills")


def test_docx_from_bytes_with_filename():
    r = convert_to_markdown(
        (FIXTURES / "sample.docx").read_bytes(), filename="Jordan Rivera Resume.DOCX"
    )
    assert r.source_format == "docx"
    assert BULLETS[0] in r.markdown  # extension match is case-insensitive


def test_pdf_converts_with_bullets_dehyphenation_and_a_structure_warning():
    r = convert_to_markdown(FIXTURES / "sample.pdf")
    assert r.source_format == "pdf"
    for b in BULLETS:
        assert f"- {b}" in r.markdown  # bullet glyph at line start -> "- "
    # "manage-\nment": "management" appears nowhere else in the fixture, so the join is not
    # safe (it could be "cross-\nfunctional"); the hyphen is kept, the newline is not, and the
    # operator is warned. The join happens when the word is seen elsewhere (tested below).
    assert "vendor manage-ment, budgeting" in r.markdown
    assert "manage-\n" not in r.markdown and "management" not in r.markdown
    assert "Jordan Rivera" in r.markdown and "Cascade Regional Health" in r.markdown
    assert any("loses structure" in w for w in r.warnings)  # the operator is told to review
    assert any("hyphen was kept" in w for w in r.warnings)


def test_hyphen_line_breaks_join_only_words_the_document_uses_unhyphenated():
    """Hard rule 2 in the importer: a joined word must be one the operator wrote. "manage-\nment"
    joins when "management" occurs elsewhere; "cross-\nfunctional" keeps its hyphen (the
    document never says "crossfunctional"); an uppercase continuation is never touched."""
    from kestrel.docs.convert import _clean_pdf_page, _join_hyphen_breaks

    pages = [
        "Vendor manage-\nment and a cross-\nfunctional team; cloud-\nNative too.",
        "Management is the word on page two.",
    ]
    vocabulary = {w.lower() for w in convert._WORD_RE.findall("\n".join(pages))}
    text, kept = _join_hyphen_breaks(pages[0], vocabulary)
    assert "Vendor management and" in text
    assert "cross-functional team" in text and "crossfunctional" not in text
    assert "cloud-\nNative" in text  # uppercase continuation: a real break, not a soft hyphen
    assert kept is True
    text, kept = _join_hyphen_breaks("Manage-\nment only, said twice: management.", {"management"})
    assert text.startswith("Management only") and kept is False
    # the page cleaner reports the kept flag so _convert_pdf can warn exactly once
    cleaned, kept = _clean_pdf_page("a-\nb here", {"ab"})
    assert cleaned == "ab here" and kept is False
    cleaned, kept = _clean_pdf_page("a-\nb here", set())
    assert cleaned == "a-b here" and kept is True


def test_pdf_warns_only_when_a_hyphen_was_kept(monkeypatch):
    class _Page:
        def __init__(self, text):
            self._t = text

        def extract_text(self):
            return self._t

    class _Reader:
        is_encrypted = False
        pages_text: list[str] = []

        def __init__(self, _fh):
            self.pages = [_Page(t) for t in self.pages_text]

    import pypdf

    monkeypatch.setattr(pypdf, "PdfReader", _Reader)
    _Reader.pages_text = ["vendor manage-\nment", "management elsewhere"]
    r = convert_to_markdown(b"%PDF-fake", filename="x.pdf")
    assert "vendor management" in r.markdown
    assert not any("hyphen was kept" in w for w in r.warnings)
    _Reader.pages_text = ["a cross-\nfunctional team"]
    r = convert_to_markdown(b"%PDF-fake", filename="x.pdf")
    assert "cross-functional team" in r.markdown
    assert sum("hyphen was kept" in w for w in r.warnings) == 1


def test_docx_merged_cells_are_read_once_and_nested_tables_warn():
    """python-docx yields a merged cell once per grid column it spans; the text must appear once
    (the operator's wording is never duplicated), and a nested table is flagged for review."""
    import docx

    d = docx.Document()
    t = d.add_table(rows=2, cols=2)
    a, b = t.rows[0].cells
    a.merge(b).text = "Skills"
    t.rows[1].cells[0].text = "Security"
    t.rows[1].cells[1].text = "Infrastructure"
    buf = io.BytesIO()
    d.save(buf)
    r = convert_to_markdown(buf.getvalue(), filename="merged.docx")
    assert "| Skills |  |" in r.markdown
    assert "| Skills | Skills |" not in r.markdown
    assert "| Security | Infrastructure |" in r.markdown
    assert r.warnings == []
    inner = t.rows[1].cells[0].add_table(rows=1, cols=1)
    inner.rows[0].cells[0].text = "nested"
    buf = io.BytesIO()
    d.save(buf)
    r = convert_to_markdown(buf.getvalue(), filename="nested.docx")
    assert "| Security | Infrastructure |" in r.markdown
    assert "nested" not in r.markdown  # python-docx cell.text omits it: the warning says so
    assert sum("nested table" in w and "missing" in w for w in r.warnings) == 1


def test_markdown_and_text_pass_through_unchanged():
    src = (FIXTURES / "sample.md").read_text(encoding="utf-8")
    r = convert_to_markdown(FIXTURES / "sample.md")
    assert r.source_format == "md"
    assert r.markdown == src  # byte-for-byte passthrough
    t = convert_to_markdown(b"\xef\xbb\xbfLine one\r\n- " + BULLETS[1].encode(), filename="a.txt")
    assert t.source_format == "txt"
    assert t.markdown == f"Line one\n- {BULLETS[1]}"  # BOM stripped, CRLF normalized


def test_every_bullet_survives_verbatim_for_the_import_check():
    """The point of the converter: a bullet P7 copies out of the converted text passes the same
    verbatim check accept_import enforces, for every format."""
    for name in ("sample.docx", "sample.pdf", "sample.md"):
        md = convert_to_markdown(FIXTURES / name).markdown
        for b in BULLETS:
            assert _verbatim(b, md), (name, b)


# ----------------------------------------------------------------------------- legacy .doc
def test_doc_without_libreoffice_is_unsupported_with_a_helpful_message(monkeypatch):
    monkeypatch.setattr(convert.shutil, "which", lambda _name: None)
    with pytest.raises(UnsupportedFormat) as ei:
        convert_to_markdown(b"\xd0\xcf\x11\xe0 fake", filename="old.doc")
    msg = str(ei.value)
    assert ".docx" in msg and ".pdf" in msg and "LibreOffice" in msg
    assert isinstance(ei.value, ConvertError)


@pytest.mark.skipif(sys.platform == "win32", reason="stand-in soffice is a POSIX shell script")
def test_doc_converts_through_soffice_in_a_temp_dir(monkeypatch, tmp_path):
    """The LibreOffice path is wired as `--headless --convert-to docx --outdir <tmp>`; a stand-in
    soffice that writes the docx fixture where the real one would proves the wiring, the outdir
    pick-up, and the source_format/warning, without needing Writer installed in CI."""
    stand_in = tmp_path / "soffice"
    stand_in.write_text(
        "#!/bin/sh\n"
        "# args: -env:... --headless --norestore --convert-to docx --outdir OUT SRC\n"
        'while [ "$1" != "--outdir" ]; do shift; done\n'
        f'cp "{FIXTURES / "sample.docx"}" "$2/input.docx"\n'
    )
    stand_in.chmod(stand_in.stat().st_mode | stat.S_IXUSR)
    monkeypatch.setattr(convert.shutil, "which", lambda _name: str(stand_in))
    r = convert_to_markdown(b"\xd0\xcf\x11\xe0 fake", filename="old.doc")
    assert r.source_format == "doc"
    assert BULLETS[0] in r.markdown
    assert any("LibreOffice" in w for w in r.warnings)


@pytest.mark.skipif(sys.platform == "win32", reason="stand-in soffice is a POSIX shell script")
def test_doc_soffice_failure_is_a_message_not_a_crash(monkeypatch, tmp_path):
    stand_in = tmp_path / "soffice"
    stand_in.write_text("#!/bin/sh\necho 'Error: source file could not be loaded' >&2\nexit 1\n")
    stand_in.chmod(stand_in.stat().st_mode | stat.S_IXUSR)
    monkeypatch.setattr(convert.shutil, "which", lambda _name: str(stand_in))
    with pytest.raises(ConvertError) as ei:
        convert_to_markdown(b"junk", filename="old.doc")
    # The converter's own stderr can name temp paths: it goes to the log, not into the message.
    assert "could not be loaded" not in str(ei.value)
    assert "LibreOffice could not convert 'old.doc'" in str(ei.value)


# ---------------------------------------------------------------------------------- refusals
def test_oversize_is_rejected_before_parsing():
    with pytest.raises(FileTooLarge):
        convert_to_markdown(b"x" * (MAX_BYTES + 1), filename="big.md")
    with pytest.raises(FileTooLarge):
        convert_to_markdown(b"x" * (MAX_BYTES + 1), filename="big.docx")  # never reaches the parser


def test_oversize_path_is_rejected(tmp_path):
    big = tmp_path / "big.txt"
    with big.open("wb") as fh:
        fh.truncate(MAX_BYTES + 1)
    with pytest.raises(FileTooLarge):
        convert_to_markdown(big)


def test_unknown_extension_is_unsupported():
    for name in ("resume.rtf", "resume.odt", "resume.html", "resume", "resume.exe"):
        with pytest.raises(UnsupportedFormat):
            convert_to_markdown(b"anything", filename=name)
    with pytest.raises(UnsupportedFormat):
        convert_to_markdown(b"anything", filename="")  # bytes need a filename


def test_scanned_pdf_with_no_text_is_a_clear_error():
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    w.write(buf)
    with pytest.raises(NoExtractableText) as ei:
        convert_to_markdown(buf.getvalue(), filename="scan.pdf")
    assert "scanned" in str(ei.value) and "OCR" in str(ei.value)


def test_corrupt_docx_is_a_message_not_a_traceback():
    with pytest.raises(ConvertError):
        convert_to_markdown(b"this is not a zip", filename="bad.docx")


# ------------------------------------------------------------------------------- web route
_NUMBERED = re.compile(r"^\[(\d+)\] (.*)$", re.M)


def _p7_v2_answer(sent: str, schema_name: str) -> dict:
    """What a faithful model answers in P7 v2 (D-IMPORT-3): line numbers only. The outline makes a
    role of every heading that has bullets under it; the role map makes each bullet line a slot."""
    rows = [(int(n), t) for n, t in _NUMBERED.findall(sent)]
    if "outline" in schema_name:
        heads = [i for i, (_, t) in enumerate(rows) if t.lstrip().startswith("#")]
        roles = []
        for k, h in enumerate(heads):
            end = heads[k + 1] if k + 1 < len(heads) else len(rows)
            seg = rows[h:end]
            if any(t.lstrip().startswith(("- ", "* ")) for _, t in seg[1:]):
                roles.append({"header_lines": [seg[0][0]], "first_line": seg[0][0],
                              "last_line": seg[-1][0], "company_label": seg[0][1].strip("# "),
                              "title": "", "start_date": "", "end_date": ""})
        return {"summary_blocks": [], "roles": roles, "skill_groups": []}
    bullets = [{"key": f"b{n}", "lines": [n], "optional": False, "tags": []}
               for n, t in rows if t.lstrip().startswith(("- ", "* "))]
    return {"context_lines": [], "bullets": bullets, "proof_points": []}


def _p7_payload_from(markdown: str) -> dict:
    """A pydantically valid P7 proposal whose derived:false variants are verbatim from the source
    (so the flags come back clean), built the way the real model is asked to."""
    assert all(b in markdown for b in BULLETS)
    return {
        "roles": [
            {
                "company_label": "Northwind Freight", "title": "VP of Information Technology",
                "start_date": "2019", "end_date": "2024", "context_line": "",
                "slots": [
                    {"key": "nw_b1", "optional": False, "default_variant_key": "v",
                     "variants": [{"key": "v", "text": BULLETS[0], "tags": [], "derived": False}]},
                    {"key": "nw_b2", "optional": False, "default_variant_key": "v",
                     "variants": [{"key": "v", "text": BULLETS[1], "tags": [], "derived": False}]},
                ],
            },
            {
                "company_label": "Cascade Regional Health", "title": "Director of Infrastructure",
                "start_date": "2014", "end_date": "2019", "context_line": "",
                "slots": [
                    {"key": "cr_b1", "optional": False, "default_variant_key": "v",
                     "variants": [{"key": "v", "text": BULLETS[2], "tags": [], "derived": False}]},
                ],
            },
        ],
        "summary_variants": [],
        "skill_groups": [],
        "proof_points": [],
    }


class _FakeP7LLM:
    """Stands in for kestrel.llm.client.LLMClient in the import route: returns a P7 proposal whose
    bullets are copied from the markdown it was sent (the user message), never invented."""

    calls: list[dict] = []

    def __init__(self, settings, http=None, **_):
        self.settings = settings

    async def complete(self, *, tier, messages, schema=None, schema_name="", **_):
        from kestrel.llm.client import ChatResult

        sent = messages[-1]["content"]
        model = self.settings.llm.tiers[tier].model
        self.calls.append({"tier": tier, "user": sent, "model": model, "schema": schema_name})
        return ChatResult(content=json.dumps(_p7_v2_answer(sent, schema_name)), model="fake")


def _import_client(monkeypatch):
    import asyncio

    from starlette.testclient import TestClient

    from kestrel import paths
    from kestrel.db import migrate
    from kestrel.db.engine import reset_engine
    from kestrel.web.app import build_app

    monkeypatch.setattr("kestrel.llm.client.LLMClient", _FakeP7LLM)
    _FakeP7LLM.calls.clear()
    migrate.run_migrations(paths.db_path())
    asyncio.run(reset_engine())
    return TestClient(build_app(worker=False))


_HOSTED_ROWS = {
    "llm.hosted_mode.enabled": True,
    "llm.hosted_mode.acknowledged_at": "2026-10-12T10:00:00+00:00",
    "llm.hosted_mode.acknowledged_terms": _TERMS,
    "llm.mode": "hosted",
    "llm.hosted_tiers.tier1.backend": "frontier",
    "llm.hosted_tiers.tier1.model": "hosted/t1",
    "llm.hosted_tiers.tier2.backend": "frontier",
    "llm.hosted_tiers.tier2.model": "hosted/t2",
    "llm.hosted_tiers.docsel.backend": "frontier",
    "llm.hosted_tiers.docsel.model": "hosted/p7-model",
}


def _seed_hosted_rows() -> None:
    """A hosted-only install: the rows `kestrel setup --hosted` writes for tier1, tier2 and
    docsel (the P7 tier, D-IMPORT-4)."""
    import asyncio

    from kestrel.db.engine import get_sessionmaker
    from kestrel.db.models import Setting

    async def _go():
        async with get_sessionmaker()() as s:
            for k, v in _HOSTED_ROWS.items():
                s.add(Setting(key=k, value=v))
            await s.commit()

    asyncio.run(_go())


def test_web_import_accepts_a_docx_upload_and_shows_extracted_text_and_proposal(monkeypatch):
    client = _import_client(monkeypatch)
    with (FIXTURES / "sample.docx").open("rb") as fh:
        r = client.post(
            "/api/documents/import",
            data={"track": "employment", "doc_kind": "resume", "text": ""},
            files={"file": ("sample.docx", fh,
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
    assert r.status_code == 200
    html = r.text
    assert 'data-testid="import-extracted"' in html  # what was read is shown first
    assert "sample.docx" in html and "(docx)" in html
    for b in BULLETS:
        assert f"- {b}" in html  # the extracted markdown, visible for review
    assert "Proposed import" in html
    assert "Every bullet is verbatim" in html  # flags clean: the proposal copied the source
    assert "Accept and create master" in html
    assert 'name="proposal_json"' in html
    assert html.index("import-extracted") < html.index("Proposed import")
    # The model was sent the CONVERTED markdown, not the (empty) textarea.
    assert _FakeP7LLM.calls and BULLETS[0] in _FakeP7LLM.calls[0]["user"]


def test_web_import_uses_the_setting_table_layer_so_hosted_rows_resolve(monkeypatch):
    """A hosted-only install keeps its models in the setting table. The route used to build its
    client on the YAML layer alone, so P7 resolved the absent local model; it must resolve the
    hosted row (and, in local mode, the YAML model)."""
    from kestrel.config import get_settings

    client = _import_client(monkeypatch)
    md = (FIXTURES / "sample.md").read_text(encoding="utf-8")
    r = client.post("/api/documents/import", data={"track": "employment", "text": md})
    assert r.status_code == 200 and "Proposed import" in r.text
    assert _FakeP7LLM.calls[-1]["tier"] == "docsel"
    assert _FakeP7LLM.calls[-1]["model"] == get_settings().llm.tiers["docsel"].model
    _seed_hosted_rows()
    r = client.post("/api/documents/import", data={"track": "employment", "text": md})
    assert r.status_code == 200 and 'data-testid="doc-online-send"' in r.text
    from kestrel.db.engine import get_sessionmaker
    from kestrel.services.online_consent import destination_id
    from kestrel.services.setting_overrides import load_settings_with_db

    async def route_dest():
        async with get_sessionmaker()() as s:
            settings = await load_settings_with_db(s)
            return settings, destination_id(settings, "docsel")

    settings, dest = asyncio.run(route_dest())
    r = client.post("/api/documents/import", data={
        "track": "employment", "text": md, "confirmed": "1",
        "model": settings.llm.tiers["docsel"].model, "dest": dest,
    })
    assert r.status_code == 200 and "Proposed import" in r.text
    assert _FakeP7LLM.calls[-1]["model"] == "hosted/p7-model"


def test_cli_doc_import_uses_the_setting_table_layer(monkeypatch, tmp_path):
    from typer.testing import CliRunner

    from kestrel.cli import app

    monkeypatch.setattr("kestrel.llm.client.LLMClient", _FakeP7LLM)
    _FakeP7LLM.calls.clear()
    _import_client(monkeypatch)  # migrates the per-test database
    _seed_hosted_rows()
    src = tmp_path / "resume.md"
    src.write_text((FIXTURES / "sample.md").read_text(encoding="utf-8"), encoding="utf-8")
    result = CliRunner().invoke(app, ["doc", "import", str(src)])
    assert result.exit_code == 0, result.output
    assert "proposed" in result.output
    assert _FakeP7LLM.calls and _FakeP7LLM.calls[-1]["model"] == "hosted/p7-model"


def test_web_import_converts_off_the_event_loop_and_bounds_the_upload(monkeypatch):
    """pypdf/python-docx parsing and the 120 s LibreOffice subprocess are synchronous: the route
    must run the converter in a thread (no running loop in the converter's thread), and an
    upload past convert.MAX_BYTES is refused before it is buffered or parsed."""
    import asyncio

    seen: dict = {}

    def _slow_convert(data, *, filename=None):
        try:
            asyncio.get_running_loop()
            seen["on_loop"] = True
        except RuntimeError:
            seen["on_loop"] = False
        seen["size"] = len(data)
        md = (FIXTURES / "sample.md").read_text(encoding="utf-8")
        return ConvertResult(markdown=md, source_format="md")

    monkeypatch.setattr(convert, "convert_to_markdown", _slow_convert)
    client = _import_client(monkeypatch)
    r = client.post("/api/documents/import", data={"track": "employment"},
                    files={"file": ("resume.md", io.BytesIO(b"# hi\n" * 10), "text/markdown")})
    assert r.status_code == 200 and "Proposed import" in r.text
    assert seen == {"on_loop": False, "size": 50}  # ran in a thread, with the whole body

    monkeypatch.setattr(convert, "MAX_BYTES", 64)
    seen.clear()
    _FakeP7LLM.calls.clear()
    big = io.BytesIO(b"x" * 65)
    r = client.post("/api/documents/import", data={"track": "employment"},
                    files={"file": ("big.md", big, "text/markdown")})
    assert r.status_code == 200 and 'class="warnbox"' in r.text
    assert "over the" in r.text and "import limit" in r.text
    assert seen == {} and not _FakeP7LLM.calls  # never converted, never sent to the model
    r = client.post("/api/documents/import", data={"track": "employment"},
                    files={"file": ("ok.md", io.BytesIO(b"x" * 64), "text/markdown")})
    assert r.status_code == 200 and seen.get("size") == 64  # at the limit is fine


def test_web_import_pdf_upload_surfaces_the_structure_warning(monkeypatch):
    client = _import_client(monkeypatch)
    with (FIXTURES / "sample.pdf").open("rb") as fh:
        r = client.post("/api/documents/import", data={"track": "employment"},
                        files={"file": ("sample.pdf", fh, "application/pdf")})
    assert r.status_code == 200
    assert "loses structure" in r.text
    assert "Proposed import" in r.text


def test_web_import_conversion_errors_render_as_a_warnbox_never_a_500(monkeypatch):
    client = _import_client(monkeypatch)
    # Unsupported extension.
    r = client.post("/api/documents/import", data={"track": "employment"},
                    files={"file": ("resume.rtf", io.BytesIO(b"{\\rtf1 hi}"), "text/rtf")})
    assert r.status_code == 200
    assert 'class="warnbox"' in r.text and "unsupported file type" in r.text
    assert not _FakeP7LLM.calls  # nothing was sent to the model
    # Corrupt docx.
    r = client.post("/api/documents/import", data={"track": "employment"},
                    files={"file": ("bad.docx", io.BytesIO(b"not a zip"), "application/zip")})
    assert r.status_code == 200 and 'class="warnbox"' in r.text
    # .doc without LibreOffice.
    monkeypatch.setattr(convert.shutil, "which", lambda _name: None)
    r = client.post("/api/documents/import", data={"track": "employment"},
                    files={"file": ("old.doc", io.BytesIO(b"\xd0\xcf\x11\xe0"),
                                    "application/msword")})
    assert r.status_code == 200 and "LibreOffice" in r.text
    # Nothing at all.
    r = client.post("/api/documents/import", data={"track": "employment", "text": ""})
    assert r.status_code == 200 and "choose a file" in r.text
    assert not _FakeP7LLM.calls


def test_web_import_pasted_markdown_still_works_without_a_file(monkeypatch):
    client = _import_client(monkeypatch)
    md = (FIXTURES / "sample.md").read_text(encoding="utf-8")
    r = client.post("/api/documents/import", data={"track": "employment", "text": md})
    assert r.status_code == 200
    assert "Proposed import" in r.text
    assert 'data-testid="import-extracted"' not in r.text  # no file: nothing to show as "read"


def test_documents_page_has_the_file_input():
    import asyncio

    from starlette.testclient import TestClient

    from kestrel.db import migrate
    from kestrel.db.engine import reset_engine
    from kestrel.web.app import build_app

    migrate.run_migrations()  # the page lists doc_masters, so it needs the schema
    asyncio.run(reset_engine())
    html = TestClient(build_app(worker=False)).get("/documents").text
    assert 'type="file"' in html and 'name="file"' in html
    assert 'accept=".md,.txt,.pdf,.docx,.doc"' in html
    assert 'hx-encoding="multipart/form-data"' in html


def test_fixtures_are_fictional():
    """No operator data in committed fixtures: the person and both organizations are invented."""
    for name in ("sample.docx", "sample.pdf", "sample.md"):
        md = convert_to_markdown(FIXTURES / name).markdown
        assert "Jordan Rivera" in md
        assert "Northwind Freight" in md and "Cascade Regional Health" in md
    assert os.path.getsize(FIXTURES / "sample.docx") < 100_000
